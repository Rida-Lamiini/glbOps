"""Batch consultation: a list of parcels in, an Excel workbook with the nearest projets of each out."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .dataset import load_entries
from .engine import consult

STATUS = {"ok": "OK", "warning": "À vérifier", "danger": "Déjà livré"}
RELATION = {"dans": "Dans", "chevauche": "Chevauche", "mitoyen": "Mitoyen", "proche": "Proche", "titre": "Même titre"}
MAX_PARCELS = 500
PER_PARCEL = 10


def run_batch(user, parcels, radius_m):
    """[(parcel, consult_result)] for up to MAX_PARCELS parcels, sharing one dataset load."""
    entries = load_entries(user)
    out = []
    for p in parcels[:MAX_PARCELS]:
        out.append((p, consult(user, ring=p.ring if p.kind == "polygon" else None, lat=p.bornes[0]["lat"], lng=p.bornes[0]["lng"],
                               titre=p.titre or None, radius_m=radius_m, entries=entries)))
    return out


def _style(ws, widths):
    head = PatternFill("solid", fgColor="1F3A2E")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_workbook(results, radius_m, source=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "Parcelles"
    ws.append(["Parcelle", "Titre", "Latitude", "Longitude", "Surface m²", "Projets proches", "Plus proche (m)", "Chevauchements", "Verdict"])
    for p, r in results:
        q, s = r["query"], r["summary"]
        ws.append([p.name, p.titre, round(q["lat"], 6), round(q["lng"], 6), q["area_m2"] or None, s["count"], s["nearest_m"], s["overlapping"], STATUS[s["status"]]])
    _style(ws, [22, 18, 12, 12, 12, 14, 14, 15, 14])
    red = PatternFill("solid", fgColor="F8D7D2")
    amber = PatternFill("solid", fgColor="FBE9C8")
    for row in ws.iter_rows(min_row=2):
        fill = red if row[8].value == STATUS["danger"] else amber if row[8].value == STATUS["warning"] else None
        if fill:
            for c in row:
                c.fill = fill

    ws2 = wb.create_sheet("Résultats")
    ws2.append(["Parcelle", "Titre consulté", "Projet voisin", "Client", "Situation", "Relation", "Distance (m)", "Direction", "Recouvrement m²", "Livré"])
    for p, r in results:
        for n in r["neighbours"][:PER_PARCEL]:
            ws2.append([p.name, p.titre, n["label"], n["client"], n["situation"], RELATION.get(n["relation"], n["relation"]),
                        n["distance_m"], n["direction"], n["overlap_m2"] or None, "oui" if n["delivered"] else ""])
    _style(ws2, [22, 18, 26, 26, 30, 12, 12, 10, 16, 8])

    ws3 = wb.create_sheet("Paramètres")
    for row in [("Fichier", source), ("Rayon de recherche (m)", radius_m), ("Parcelles traitées", len(results)),
                ("Voisins listés par parcelle", f"{PER_PARCEL} (les plus proches)")]:
        ws3.append(list(row))
    ws3.column_dimensions["A"].width, ws3.column_dimensions["B"].width = 30, 40
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

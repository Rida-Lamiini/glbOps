"""Reads a file (or pasted text) describing a parcel or a set of points into `Parcel` objects.

Accepted: CSV / TXT / TSV / Excel (.xlsx) with bornes (name, X, Y in Lambert — or lat/lng), the text copied from the
ANCFCC "consultation de la mappe cadastrale" page (``B14 X : 369020.33 Y : 371666.77``), KML, GeoJSON and GPX.
The coordinate system is detected per point: large X/Y values are Lambert (Nord or Sud, given by the caller),
small ones are WGS84 latitude/longitude.
"""

import csv
import io
import json
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field

from cadastre.geo.proj import lambert_to_wgs84

NUM = r"-?\d+(?:[.,]\d+)?"
BORNE_RE = re.compile(rf"(?P<name>[A-Za-z]{{1,4}}\.?\s?\d+[A-Za-z]?)\s*X\s*[:=]?\s*(?P<x>{NUM})\s*Y\s*[:=]?\s*(?P<y>{NUM})", re.I)
TITRE_RE = re.compile(r"R[ée]f[ée]rence\s+fonci[èe]re\s*:?\s*(?P<t>[A-Za-z]{0,3}\s?\d[\w/\-]*)", re.I)
SURFACE_RE = re.compile(r"(?:(?P<ha>\d+)\s*ha\s*)?(?P<a>\d+)\s*a\s*(?P<ca>\d+)\s*ca", re.I)
HEADERS = {
    "name": {"borne", "bornes", "point", "nom", "name", "n", "label", "designation"},
    "parcel": {"parcelle", "parcel", "titre", "titre foncier", "tf", "reference", "ref", "dossier", "reference fonciere", "lot"},
    "x": {"x", "x lambert", "easting", "est"},
    "y": {"y", "y lambert", "northing", "nord"},
    "lat": {"lat", "latitude"},
    "lng": {"lng", "lon", "long", "longitude"},
}


@dataclass
class Parcel:
    name: str
    titre: str = ""
    kind: str = "polygon"  # point | line | polygon
    bornes: list = field(default_factory=list)  # [{name, x, y, lat, lng}] — x/y are None for lat/lng input
    zone: str = "nord"
    declared_surface_m2: float | None = None
    source: str = ""

    @property
    def ring(self):
        return [(b["lat"], b["lng"]) for b in self.bornes]

    def to_dict(self):
        return {
            "name": self.name, "titre": self.titre, "kind": self.kind, "zone": self.zone,
            "declared_surface_m2": self.declared_surface_m2, "source": self.source,
            "bornes": self.bornes,
        }


class ParseError(ValueError):
    pass


def _num(value):
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", "").replace(" ", "")
    if re.fullmatch(NUM, s):
        return float(s.replace(",", "."))
    return None


def _norm(s):
    return re.sub(r"[^a-z0-9 ]+", "", str(s).lower().replace("é", "e").replace("è", "e")).strip()


def _classify(a, b):
    """('lambert', x, y) | ('wgs84', lat, lng) | None for a pair of numbers read in the order they were given."""
    if abs(a) > 1000 and abs(b) > 1000:
        return "lambert", a, b
    if 20 <= a <= 40 and -20 <= b <= 0:
        return "wgs84", a, b
    if -20 <= a <= 0 and 20 <= b <= 40:
        return "wgs84", b, a  # written longitude, latitude
    if abs(a) <= 90 and abs(b) <= 180:
        return "wgs84", a, b
    return None


def _borne(name, kind, a, b, zone):
    if kind == "lambert":
        lat, lng = lambert_to_wgs84(a, b, zone)
        return {"name": name, "x": a, "y": b, "lat": lat, "lng": lng}
    return {"name": name, "x": None, "y": None, "lat": a, "lng": b}


def _dedupe(bornes):
    out = []
    for b in bornes:
        if out and (out[-1]["lat"], out[-1]["lng"]) == (b["lat"], b["lng"]):
            continue
        out.append(b)
    if len(out) > 1 and (out[0]["lat"], out[0]["lng"]) == (out[-1]["lat"], out[-1]["lng"]):
        out.pop()  # the contour is closed by repeating the first point: not a borne
    return out


def _parcel(name, titre, bornes, zone, mode, source, surface=None):
    bornes = _dedupe(bornes)
    if not bornes:
        return None
    if mode == "points":
        kind = "point"
    elif mode == "polygon":
        kind = "polygon" if len(bornes) >= 3 else ("line" if len(bornes) == 2 else "point")
    else:  # auto
        kind = "polygon" if len(bornes) >= 3 else ("line" if len(bornes) == 2 else "point")
    return Parcel(name=name, titre=titre, kind=kind, bornes=bornes, zone=zone, declared_surface_m2=surface, source=source)


def _split_points(parcel_name, titre, bornes, zone, mode, source, surface):
    """`points` mode: every borne is its own parcel."""
    return [
        Parcel(name=b["name"], titre=titre, kind="point", bornes=[b], zone=zone, source=source)
        for b in _dedupe(bornes)
    ]


def _surface_m2(text):
    m = SURFACE_RE.search(text)
    if not m:
        return None
    return int(m.group("ha") or 0) * 10000 + int(m.group("a")) * 100 + int(m.group("ca"))


# ------------------------------------------------------------------------------------------------------ text / tables
def _decode(data):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _parse_ancfcc_text(text, zone, mode, source):
    found = list(BORNE_RE.finditer(text))
    if not found:
        return None
    bornes = []
    for m in found:
        x, y = _num(m.group("x")), _num(m.group("y"))
        c = _classify(x, y)
        if c:
            bornes.append(_borne(re.sub(r"\s+", "", m.group("name")).upper(), c[0], c[1], c[2], zone))
    t = TITRE_RE.search(text)
    titre = re.sub(r"\s+", "", t.group("t")) if t else ""
    surface = _surface_m2(text)
    if mode == "points":
        return _split_points(titre or "Points", titre, bornes, zone, mode, source, surface)
    p = _parcel(titre or "Parcelle", titre, bornes, zone, mode, source, surface)
    return [p] if p else None


def _rows_from_text(text):
    sample = "\n".join(text.splitlines()[:20])
    if ";" in sample:
        delim = ";"
    elif "\t" in sample:
        delim = "\t"
    elif all("," in l for l in sample.splitlines() if l.strip()) and not re.search(rf"{NUM}[ 	]+{NUM}", sample):
        delim = ","
    else:
        delim = None  # whitespace
    rows = []
    for line in text.splitlines():
        if not line.strip():
            continue
        rows.append([c.strip() for c in (next(csv.reader([line], delimiter=delim)) if delim else line.split())])
    return rows


def _header_roles(row):
    roles = {}
    for i, cell in enumerate(row):
        if isinstance(cell, (int, float)) or _num(cell) is not None:
            return {}
        key = _norm(cell)
        for role, words in HEADERS.items():
            if key in words and role not in roles:
                roles[role] = i
    return roles if ("x" in roles and "y" in roles) or ("lat" in roles and "lng" in roles) else {}


def _from_rows(rows, zone, mode, source):
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []
    roles = _header_roles(rows[0])
    body = rows[1:] if roles else rows
    records = []
    for n, row in enumerate(body, start=1):
        row = list(row)
        if roles:
            get = lambda role: row[roles[role]] if role in roles and roles[role] < len(row) else ""
            if "x" in roles:
                a, b, order = _num(get("x")), _num(get("y")), "xy"
            else:
                a, b, order = _num(get("lat")), _num(get("lng")), "latlng"
            name = str(get("name")).strip() or f"P{n}"
            key = str(get("parcel")).strip()
        else:
            nums = [(i, _num(c)) for i, c in enumerate(row) if _num(c) is not None]
            if len(nums) < 2:
                continue
            # the two last numeric columns hold the coordinates (a leading number is usually a sequence no.)
            (_, a), (_, b) = nums[-2], nums[-1]
            labels = [str(c).strip() for c in row if _num(c) is None and str(c).strip()]
            name = labels[0] if labels else f"P{n}"
            key, order = "", None
        if a is None or b is None:
            continue
        c = _classify(a, b)
        if not c:
            continue
        records.append({"key": key, "bornes": _borne(name, c[0], c[1], c[2], zone)})
    groups = {}
    for r in records:
        groups.setdefault(r["key"], []).append(r["bornes"])
    parcels = []
    for key, bornes in groups.items():
        if mode == "points" or (mode == "auto" and key and len(bornes) < 3 and False):
            parcels.extend(_split_points(key or "Points", key, bornes, zone, mode, source, None))
        else:
            p = _parcel(key or "Parcelle", key, bornes, zone, mode, source)
            if p:
                parcels.append(p)
    return parcels


def _from_xlsx(data, zone, mode, source):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parcels = []
    for ws in wb.worksheets:
        rows = [["" if c is None else c for c in row] for row in ws.iter_rows(values_only=True)]
        parcels.extend(_from_rows(rows, zone, mode, source))
        if parcels:
            break
    return parcels


# ------------------------------------------------------------------------------------------------------ geo formats
def _local(tag):
    return tag.rsplit("}", 1)[-1]


def _coords_text(text):
    out = []
    for token in text.split():
        parts = token.split(",")
        if len(parts) >= 2:
            lng, lat = _num(parts[0]), _num(parts[1])
            if lng is not None and lat is not None:
                out.append((lat, lng))
    return out


def _bornes_from_latlng(pairs, prefix="P"):
    return [{"name": f"{prefix}{i}", "x": None, "y": None, "lat": lat, "lng": lng} for i, (lat, lng) in enumerate(pairs, start=1)]


def _from_kml(data, zone, mode, source):
    root = ET.fromstring(data)
    parcels = []
    for pm in (e for e in root.iter() if _local(e.tag) == "Placemark"):
        name = next((e.text.strip() for e in pm if _local(e.tag) == "name" and e.text), "")
        for geom in pm.iter():
            kind = _local(geom.tag)
            if kind not in ("Polygon", "LineString", "Point"):
                continue
            coords = next((e for e in geom.iter() if _local(e.tag) == "coordinates"), None)
            pairs = _coords_text(coords.text or "") if coords is not None else []
            if kind == "Polygon":  # only the outer boundary
                outer = next((e for e in geom.iter() if _local(e.tag) == "outerBoundaryIs"), None)
                if outer is not None:
                    c = next((e for e in outer.iter() if _local(e.tag) == "coordinates"), None)
                    pairs = _coords_text(c.text or "") if c is not None else pairs
            p = _parcel(name or "Parcelle", name, _bornes_from_latlng(pairs), zone, "points" if kind == "Point" else mode, source)
            if p:
                parcels.append(p)
    return parcels


def _from_geojson(data, zone, mode, source):
    doc = json.loads(data)
    feats = doc["features"] if doc.get("type") == "FeatureCollection" else [doc] if doc.get("type") == "Feature" else [{"geometry": doc, "properties": {}}]
    parcels = []
    for f in feats:
        g = f.get("geometry") or {}
        props = f.get("properties") or {}
        name = str(props.get("titre") or props.get("name") or props.get("nom") or props.get("ref") or "Parcelle")
        t, c = g.get("type"), g.get("coordinates")
        shapes = []
        if t == "Polygon":
            shapes.append(("polygon", c[0]))
        elif t == "MultiPolygon":
            shapes.extend(("polygon", poly[0]) for poly in c)
        elif t == "LineString":
            shapes.append(("line", c))
        elif t == "Point":
            shapes.append(("point", [c]))
        for _, ring in shapes:
            pairs = [(pt[1], pt[0]) for pt in ring]
            p = _parcel(name, name if t != "Point" else "", _bornes_from_latlng(pairs), zone, "points" if t == "Point" else mode, source)
            if p:
                parcels.append(p)
    return parcels


def _from_gpx(data, zone, mode, source):
    root = ET.fromstring(data)
    pts = []
    for e in root.iter():
        if _local(e.tag) in ("wpt", "trkpt"):
            lat, lng = _num(e.get("lat", "")), _num(e.get("lon", ""))
            name = next((c.text.strip() for c in e if _local(c.tag) == "name" and c.text), "")
            if lat is not None and lng is not None:
                pts.append({"name": name or f"P{len(pts) + 1}", "x": None, "y": None, "lat": lat, "lng": lng})
    return [Parcel(name=b["name"], kind="point", bornes=[b], zone=zone, source=source) for b in pts] if mode == "points" or len(pts) < 3 else [_parcel("Tracé", "", pts, zone, mode, source)]


# ------------------------------------------------------------------------------------------------------ entry points
def parse_bytes(filename, data, zone="nord", mode="auto"):
    """Returns (parcels, warnings). Raises ParseError (French message) when nothing usable is found."""
    if zone not in ("nord", "sud"):
        zone = "nord"
    ext = (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""
    source = filename or "texte"
    try:
        if ext == "xlsx":
            parcels = _from_xlsx(data, zone, mode, source)
        elif ext == "kml":
            parcels = _from_kml(data, zone, mode, source)
        elif ext in ("geojson", "json"):
            parcels = _from_geojson(data, zone, mode, source)
        elif ext == "gpx":
            parcels = _from_gpx(data, zone, mode, source)
        else:
            parcels = parse_text(_decode(data), zone, mode, source)
    except (ET.ParseError, json.JSONDecodeError, KeyError, TypeError) as exc:
        raise ParseError(f"Fichier illisible ({exc}).") from exc
    parcels = [p for p in parcels if p]
    if not parcels:
        raise ParseError("Aucune coordonnée trouvée : attendu des bornes (nom, X, Y en Lambert — ou latitude, longitude).")
    return parcels, []


def parse_text(text, zone="nord", mode="auto", source="texte"):
    """The ANCFCC page copy first (it has the labelled `X : … Y : …` bornes), then plain rows."""
    found = _parse_ancfcc_text(text, zone, mode, source)
    if found:
        return found
    return _from_rows(_rows_from_text(text), zone, mode, source)
